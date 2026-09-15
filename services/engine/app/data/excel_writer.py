"""Append/update a single draw row in the Excel audit source.

The workbook is normally read-only (see ``excel_loader``); this module is the
one sanctioned way to extend it, used by the admin "เพิ่มงวดใหม่" flow. Every
write is preceded by a timestamped backup because the workbook is an audit
source, and openpyxl round-trips the file so existing styling survives.

Rows are stored newest-first, so a new draw is inserted before the first row
whose date is older than it (the loader sorts by date anyway — ordering here is
for humans and for the Phase-1 import script).
"""
from __future__ import annotations

import datetime
import os
import re
import shutil
from typing import Dict, List, Optional, Tuple

import openpyxl

from .excel_loader import (
    COL_B2,
    COL_DATE,
    COL_SET,
    COL_SIX,
    COL_T2,
    COL_T3,
    COL_TIME,
    HEADER_KEY,
    _find_header,
    _parse_date,
)

_DIGITS = re.compile(r"^\d+$")

# field name -> (column offset, expected digit length)
FIELD_COLS: List[Tuple[str, int, int]] = [
    ("six", COL_SIX, 6),
    ("top3", COL_T3, 3),
    ("top2", COL_T2, 2),
    ("bottom2", COL_B2, 2),
    ("set1", COL_SET[0], 3),
    ("set2", COL_SET[1], 3),
    ("set3", COL_SET[2], 3),
    ("set4", COL_SET[3], 3),
]

THAI_MONTH_ABBR = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
                   "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
BE_OFFSET = 543


class DrawWriteError(ValueError):
    """Raised when the submitted draw fails validation (never partially written)."""


def _clean(v) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def validate(payload: Dict) -> Tuple[datetime.date, Dict[str, Optional[str]], Optional[datetime.time]]:
    """Validate a draw payload. Returns (date, numbers, time) or raises."""
    raw_date = _clean(payload.get("date"))
    if not raw_date:
        raise DrawWriteError("ต้องระบุวันที่งวด")
    d = _parse_date(raw_date)
    if d is None:
        raise DrawWriteError(f"วันที่ไม่ถูกต้อง: {raw_date!r} (ต้องเป็น YYYY-MM-DD)")

    numbers: Dict[str, Optional[str]] = {}
    for name, _col, length in FIELD_COLS:
        v = _clean(payload.get(name))
        if v is not None and (not _DIGITS.match(v) or len(v) != length):
            raise DrawWriteError(f"{name} ต้องเป็นตัวเลข {length} หลัก (ได้ {v!r})")
        numbers[name] = v

    six, t3, t2 = numbers["six"], numbers["top3"], numbers["top2"]
    if six and t3 and six[-3:] != t3:
        raise DrawWriteError(f"3 บน ({t3}) ไม่ตรงกับ 3 ตัวท้ายของรางวัลที่ 1 ({six[-3:]})")
    if six and t2 and six[-2:] != t2:
        raise DrawWriteError(f"2 บน ({t2}) ไม่ตรงกับ 2 ตัวท้ายของรางวัลที่ 1 ({six[-2:]})")

    t: Optional[datetime.time] = None
    raw_time = _clean(payload.get("drawTime"))
    if raw_time:
        m = re.match(r"^(\d{1,2}):(\d{2})$", raw_time)
        if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
            raise DrawWriteError(f"เวลาไม่ถูกต้อง: {raw_time!r} (ต้องเป็น HH:MM)")
        t = datetime.time(int(m.group(1)), int(m.group(2)))

    return d, numbers, t


def _backup(path: str) -> str:
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    backup_dir = os.path.join(os.path.dirname(path) or ".", "backups")
    os.makedirs(backup_dir, exist_ok=True)
    base = os.path.basename(path)
    dest = os.path.join(backup_dir, f"{os.path.splitext(base)[0]}-{stamp}.xlsx")
    shutil.copy2(path, dest)
    return dest


def _refresh_title(ws, header_row: int) -> None:
    """Keep the A1 caption's ข้อมูล(... – ...) range in sync with the data."""
    cell = ws.cell(row=1, column=1)
    if not isinstance(cell.value, str) or "ข้อมูล(" not in cell.value:
        return
    dates = []
    for r in range(header_row + 1, ws.max_row + 1):
        d = _parse_date(ws.cell(row=r, column=COL_DATE + 1).value)
        if d is not None:
            dates.append(d)
    if not dates:
        return

    def fmt(d: datetime.date) -> str:
        return f"{d.day} {THAI_MONTH_ABBR[d.month - 1]} {d.year + BE_OFFSET}"

    cell.value = f"ข้อมูล({fmt(min(dates))} – {fmt(max(dates))})"


def upsert_draw(path: str, payload: Dict) -> Dict:
    """Insert or update one draw row. Returns a summary of what happened."""
    d, numbers, t = validate(payload)

    # The workbook stores 3 บน / 2 บน as formulas (=MID(B{r},4,3)). openpyxl
    # keeps formulas but DROPS Excel's cached results on save, which would make
    # those columns read back as None for every row (the loader uses
    # data_only=True). So we load the cached values and persist them as
    # literals — the workbook stays semantically identical and stops depending
    # on Excel having recalculated it.
    wb_formulas = openpyxl.load_workbook(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws_formulas = wb_formulas.active
    ws = wb.active
    header_row = _find_header(ws)

    # Guard: a formula whose cached value is missing means Excel never
    # recalculated this workbook. Flattening it would silently blank real data,
    # so refuse rather than corrupt the audit source.
    flattened = 0
    stale: List[str] = []
    for row in ws_formulas.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                flattened += 1
                if ws.cell(row=cell.row, column=cell.column).value is None:
                    stale.append(cell.coordinate)
    if stale:
        raise DrawWriteError(
            "ไฟล์ Excel มีสูตรที่ยังไม่ถูกคำนวณ "
            f"({len(stale)} เซลล์ เช่น {', '.join(stale[:5])}) — "
            "กรุณาเปิดไฟล์ด้วย Excel แล้วบันทึกใหม่ก่อนเพิ่มงวด"
        )

    # Locate an existing row for this date, and the insert position that keeps
    # the sheet in descending-date order.
    target: Optional[int] = None
    insert_at = ws.max_row + 1
    found_any = False
    for r in range(header_row + 1, ws.max_row + 1):
        rd = _parse_date(ws.cell(row=r, column=COL_DATE + 1).value)
        if rd is None:
            continue
        found_any = True
        if rd == d:
            target = r
            break
        if rd < d and insert_at > r:
            insert_at = r
    if not found_any:
        insert_at = header_row + 1

    action = "updated"
    before: Optional[Dict[str, Optional[str]]] = None
    if target is None:
        ws.insert_rows(insert_at)
        target = insert_at
        action = "inserted"
    else:
        before = {
            name: _clean(ws.cell(row=target, column=col + 1).value)
            for name, col, _len in FIELD_COLS
        }

    ws.cell(row=target, column=COL_DATE + 1).value = d.isoformat()
    for name, col, _len in FIELD_COLS:
        ws.cell(row=target, column=col + 1).value = numbers[name]
    ws.cell(row=target, column=COL_TIME + 1).value = t

    _refresh_title(ws, header_row)

    backup = _backup(path)
    # Write to a sibling temp file, then replace — a crash mid-save must never
    # leave a truncated audit source behind.
    tmp = f"{path}.tmp"
    wb.save(tmp)
    os.replace(tmp, path)

    return {
        "action": action,
        "formulas_flattened": flattened,
        "date": d.isoformat(),
        "row": target,
        "backup": backup,
        "before": before,
        "after": {"date": d.isoformat(), **numbers,
                  "drawTime": t.strftime("%H:%M") if t else None},
    }
